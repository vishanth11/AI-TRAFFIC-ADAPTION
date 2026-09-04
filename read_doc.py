from docx import Document

file_path = "data/raw/Traffic Flow Prediction Dataset.docx"

doc = Document(file_path)

print("=" * 70)
print("TRAFFIC FLOW DATASET DOCUMENTATION")
print("=" * 70)

for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()

    if text:
        print(f"\n[{i}] {text}")

for table_index, table in enumerate(doc.tables):

    print("\n" + "=" * 70)
    print(f"TABLE {table_index}")
    print("=" * 70)

    for row in table.rows:
        print(" | ".join(cell.text.strip() for cell in row.cells))